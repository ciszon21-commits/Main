"""
Django Management Command: import_preset_hotspots
==================================================
從 Excel 匯入 360 預設熱點資料庫，並將圖片複製到 Django media 目錄。

用法：
    python manage.py import_preset_hotspots
    python manage.py import_preset_hotspots --excel "D:/your/360預設資料庫整理.xlsx"
    python manage.py import_preset_hotspots --images-root "D:/your/images_base_dir"
    python manage.py import_preset_hotspots --dry-run
    python manage.py import_preset_hotspots --clear

資料夾對應說明：
    Excel 中「資料夾名稱」欄位的值（如 台北勞檢、桃園勞檢、職安署）
    會對應到 --images-root 目錄下的同名子資料夾。
"""

import os
import shutil
import openpyxl
from pathlib import Path
from django.core.management.base import BaseCommand, CommandError
from django.conf import settings
from site360.models import PresetHotspot, HazardType

# ============================================================
# ✏️  可自由調整的預設路徑
# ============================================================

# Excel 主資料檔案路徑
DEFAULT_EXCEL_PATH = r"E:\NickChang\02 平台開發\115 Sino360\20260225 瑞澤提供360平台資料\360預設資料庫整理.xlsx"

# 圖片資料夾的上層根目錄（該目錄下應有 台北勞檢/桃園勞檢/職安署 等子資料夾）
DEFAULT_IMAGES_ROOT = r"E:\NickChang\02 平台開發\115 Sino360\20260225 瑞澤提供360平台資料"

# Excel 工作表名稱
DEFAULT_SHEET_NAME = "工作表1"

# ============================================================


class Command(BaseCommand):
    help = "從 Excel 匯入 360 預設熱點資料庫，並複製圖片至 media 目錄"

    def add_arguments(self, parser):
        parser.add_argument(
            "--excel",
            type=str,
            default=DEFAULT_EXCEL_PATH,
            help=f"Excel 檔案路徑 (預設: {DEFAULT_EXCEL_PATH})",
        )
        parser.add_argument(
            "--images-root",
            type=str,
            default=DEFAULT_IMAGES_ROOT,
            help=f"圖片根目錄（子資料夾名稱需與 Excel 中「資料夾名稱」一致）\n(預設: {DEFAULT_IMAGES_ROOT})",
        )
        parser.add_argument(
            "--sheet",
            type=str,
            default=DEFAULT_SHEET_NAME,
            help=f"工作表名稱 (預設: {DEFAULT_SHEET_NAME})",
        )
        parser.add_argument(
            "--dry-run",
            action="store_true",
            help="僅預覽，不寫入資料庫也不複製圖片",
        )
        parser.add_argument(
            "--clear",
            action="store_true",
            help="匯入前先清空 PresetHotspot 資料表（危險操作）",
        )
        parser.add_argument(
            "--skip-images",
            action="store_true",
            help="不複製圖片，僅匯入文字資料",
        )

    def handle(self, *args, **options):
        excel_path   = options["excel"]
        images_root  = options["images_root"]
        sheet_name   = options["sheet"]
        dry_run      = options["dry_run"]
        clear        = options["clear"]
        skip_images  = options["skip_images"]

        self.stdout.write(f"\n📂 Excel 路徑   : {excel_path}")
        self.stdout.write(f"🖼️  圖片根目錄  : {images_root}")
        self.stdout.write(f"📋 工作表       : {sheet_name}\n")

        # ---- 1. 讀取 Excel ----
        try:
            wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
        except FileNotFoundError:
            raise CommandError(f"找不到 Excel 檔案：{excel_path}")
        except Exception as e:
            raise CommandError(f"讀取 Excel 失敗：{e}")

        if sheet_name not in wb.sheetnames:
            raise CommandError(
                f"找不到工作表「{sheet_name}」。可用工作表：{', '.join(wb.sheetnames)}"
            )

        ws = wb[sheet_name]
        all_rows = list(ws.iter_rows(values_only=True))
        wb.close()

        data_rows = [r for r in all_rows[1:] if any(cell for cell in r)]
        self.stdout.write(f"✅ 共讀取到 {len(data_rows)} 筆有效資料\n")

        # ---- 2. 預先建立 HazardType name → object 查找表 ----
        hazard_map = {ht.name: ht for ht in HazardType.objects.all()}

        # ---- 3. media 目標目錄 ----
        media_dest_dir = Path(settings.MEDIA_ROOT) / "site360" / "preset_hotspots"
        if not dry_run:
            media_dest_dir.mkdir(parents=True, exist_ok=True)

        # ---- 4. 解析並處理每一筆 ----
        records = []
        missing_images = []

        for row_idx, row in enumerate(data_rows, start=2):
            serial_number   = row[0]
            source_folder   = str(row[1]).strip() if row[1] else ""
            original_filename = str(row[2]).strip() if row[2] else ""
            title           = str(row[3]).strip() if row[3] else ""
            description     = str(row[4]).strip() if row[4] else ""
            hazard_type_name = str(row[5]).strip() if row[5] else ""

            # 驗證必填
            if not title:
                self.stdout.write(self.style.WARNING(f"  ⚠️  第 {row_idx} 列「標題」為空，跳過"))
                continue
            try:
                serial_number = int(serial_number)
            except (ValueError, TypeError):
                self.stdout.write(self.style.WARNING(f"  ⚠️  第 {row_idx} 列「項次」非整數（{serial_number}），跳過"))
                continue

            # 危害類型關聯
            hazard_obj = hazard_map.get(hazard_type_name)
            if hazard_type_name and not hazard_obj:
                self.stdout.write(self.style.WARNING(
                    f"  ⚠️  [{serial_number}] 危害類型「{hazard_type_name}」未找到（HazardType 未建立或名稱不符）"
                ))

            # 圖片路徑解析
            image_src_path = None
            media_relative_path = ""

            if original_filename and source_folder and not skip_images:
                src = Path(images_root) / source_folder / original_filename
                if src.exists():
                    image_src_path = src
                    # 目標：在 media 目錄下按資料夾分類儲存
                    media_relative_path = f"site360/preset_hotspots/{source_folder}/{original_filename}"
                else:
                    missing_images.append(f"  [{serial_number}] {src}")

            records.append({
                "serial_number": serial_number,
                "source_folder": source_folder,
                "original_filename": original_filename,
                "title": title,
                "description": description,
                "hazard_type": hazard_obj,
                "image_src_path": image_src_path,
                "media_relative_path": media_relative_path,
            })

        # ---- 5. 報告缺圖 ----
        if missing_images:
            self.stdout.write(self.style.WARNING(
                f"\n⚠️  以下 {len(missing_images)} 筆找不到對應圖片："
            ))
            for m in missing_images[:20]:
                self.stdout.write(f"    {m}")
            if len(missing_images) > 20:
                self.stdout.write(f"    ... 共 {len(missing_images)} 筆（只顯示前 20）")

        # ---- 6. Dry-run 預覽 ----
        if dry_run:
            self.stdout.write(f"\n📋 預覽前 10 筆資料：")
            for r in records[:10]:
                img_status = "✅ 有圖" if r["image_src_path"] else ("⬜ 無圖" if not r["original_filename"] else "❌ 圖片缺失")
                self.stdout.write(
                    f"  [{r['serial_number']:3d}][{r['source_folder']:<6}] {r['title'][:40]} "
                    f"| {r['hazard_type'].name if r['hazard_type'] else '(無危害類型)'} | {img_status}"
                )
            self.stdout.write(self.style.WARNING(f"\n⚠️  Dry-run 模式，未寫入資料庫。"))
            return

        # ---- 7. 清空（可選）----
        if clear:
            count = PresetHotspot.objects.count()
            PresetHotspot.objects.all().delete()
            self.stdout.write(self.style.WARNING(f"🗑️  已清空 PresetHotspot（刪除 {count} 筆）"))

        # ---- 8. 複製圖片 + 寫入資料庫 ----
        created_count = 0
        updated_count = 0
        image_copied  = 0
        image_failed  = 0

        for r in records:
            # 複製圖片
            image_field_value = r["media_relative_path"]
            if r["image_src_path"] and r["media_relative_path"]:
                dest = Path(settings.MEDIA_ROOT) / r["media_relative_path"]
                dest.parent.mkdir(parents=True, exist_ok=True)
                try:
                    shutil.copy2(r["image_src_path"], dest)
                    image_copied += 1
                except Exception as e:
                    self.stdout.write(self.style.ERROR(
                        f"  ❌ 複製圖片失敗 [{r['serial_number']}] {r['image_src_path']}: {e}"
                    ))
                    image_field_value = ""
                    image_failed += 1

            # 寫入資料庫（以 serial_number + source_folder 作為唯一識別）
            obj, created = PresetHotspot.objects.update_or_create(
                serial_number=r["serial_number"],
                source_folder=r["source_folder"],
                defaults={
                    "original_filename": r["original_filename"],
                    "title": r["title"],
                    "description": r["description"],
                    "hazard_type": r["hazard_type"],
                    "image": image_field_value,
                },
            )
            if created:
                created_count += 1
            else:
                updated_count += 1

        self.stdout.write(self.style.SUCCESS(
            f"\n🎉 匯入完成！\n"
            f"   資料庫：新增 {created_count} 筆 | 更新 {updated_count} 筆\n"
            f"   圖片：複製成功 {image_copied} 張 | 失敗 {image_failed} 張 | 缺失 {len(missing_images)} 張"
        ))
