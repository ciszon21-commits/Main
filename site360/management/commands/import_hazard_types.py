"""
Django Management Command: import_hazard_types
=============================================
匯入危害類型資料從 Excel 檔案到 HazardType 資料表。

用法：
    python manage.py import_hazard_types
    python manage.py import_hazard_types --path "C:/your/path/危害類型定義.xlsx"
    python manage.py import_hazard_types --sheet "工作表1" --dry-run

預設設定 (可直接在此修改)：
"""

import openpyxl
from django.core.management.base import BaseCommand, CommandError
from site360.models import HazardType

# ============================================================
# ✏️  可自由調整的預設路徑與工作表名稱
# ============================================================
DEFAULT_EXCEL_PATH = r"E:\NickChang\02 平台開發\115 Sino360\20260225 瑞澤提供360平台資料\危害類型定義.xlsx"
DEFAULT_SHEET_NAME = "工作表1"
# ============================================================


class Command(BaseCommand):
    help = "從 Excel 匯入危害類型 (HazardType) 資料"

    def add_arguments(self, parser):
        parser.add_argument(
            "--path",
            type=str,
            default=DEFAULT_EXCEL_PATH,
            help=f"Excel 檔案路徑 (預設: {DEFAULT_EXCEL_PATH})",
        )
        parser.add_argument(
            "--sheet",
            type=str,
            default=DEFAULT_SHEET_NAME,
            help=f"工作表名稱 (預設: {DEFAULT_SHEET_NAME})",
        )
        parser.add_argument(
            "--dry-run",
            action="store_true",
            help="僅預覽資料，不實際寫入資料庫",
        )
        parser.add_argument(
            "--clear",
            action="store_true",
            help="匯入前先清空 HazardType 資料表（危險操作，請謹慎）",
        )

    def handle(self, *args, **options):
        path = options["path"]
        sheet_name = options["sheet"]
        dry_run = options["dry_run"]
        clear = options["clear"]

        self.stdout.write(f"\n📂 讀取 Excel：{path}")
        self.stdout.write(f"📋 工作表：{sheet_name}\n")

        # --- 讀取 Excel ---
        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        except FileNotFoundError:
            raise CommandError(f"找不到檔案：{path}")
        except Exception as e:
            raise CommandError(f"無法讀取 Excel：{e}")

        if sheet_name not in wb.sheetnames:
            raise CommandError(
                f"找不到工作表「{sheet_name}」。\n"
                f"可用工作表：{', '.join(wb.sheetnames)}"
            )

        ws = wb[sheet_name]

        # --- 解析資料（跳過標題列）---
        records = []
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            # 空白列跳過
            if not any(cell for cell in row):
                continue

            serial_number = row[0]  # 項次
            name = row[1]           # 危害類型名稱
            description = row[2] if len(row) > 2 else ""  # 危害類型說明（可選）

            # 驗證必填欄位
            if serial_number is None or name is None:
                self.stdout.write(
                    self.style.WARNING(f"  ⚠️  第 {row_idx} 列缺少必填欄位，跳過：{row}")
                )
                continue

            try:
                serial_number = int(serial_number)
            except (ValueError, TypeError):
                self.stdout.write(
                    self.style.WARNING(f"  ⚠️  第 {row_idx} 列「項次」非整數，跳過：{serial_number}")
                )
                continue

            records.append({
                "serial_number": serial_number,
                "name": str(name).strip(),
                "description": str(description).strip() if description else "",
            })

        wb.close()

        if not records:
            raise CommandError("Excel 中沒有有效資料。")

        self.stdout.write(f"✅ 解析到 {len(records)} 筆資料：\n")

        # --- 預覽 ---
        for r in records:
            self.stdout.write(
                f"  [{r['serial_number']:3d}] {r['name']}"
                + (f"  ── {r['description'][:40]}…" if len(r.get("description", "")) > 40 else
                   (f"  ── {r['description']}" if r.get("description") else ""))
            )

        if dry_run:
            self.stdout.write(self.style.WARNING("\n⚠️  Dry-run 模式，未寫入資料庫。"))
            return

        # --- 清空（可選）---
        if clear:
            count = HazardType.objects.count()
            HazardType.objects.all().delete()
            self.stdout.write(self.style.WARNING(f"\n🗑️  已清空 HazardType 資料表（刪除 {count} 筆）"))

        # --- 寫入資料庫 ---
        created_count = 0
        updated_count = 0

        for r in records:
            obj, created = HazardType.objects.update_or_create(
                serial_number=r["serial_number"],
                defaults={
                    "name": r["name"],
                    "description": r["description"],
                },
            )
            if created:
                created_count += 1
            else:
                updated_count += 1

        self.stdout.write(
            self.style.SUCCESS(
                f"\n🎉 匯入完成！新增: {created_count} 筆 | 更新: {updated_count} 筆"
            )
        )
