import openpyxl
import json

def main():
    try:
        wb = openpyxl.load_workbook('reference/容積計算.xlsx', data_only=False)
        wb_data = openpyxl.load_workbook('reference/容積計算.xlsx', data_only=True)
        out = {}
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            ws_data = wb_data[sheet_name]
            sheet_data = []
            for r in range(1, min(ws.max_row, 100) + 1):
                row_data = []
                has_data = False
                for c in range(1, min(ws.max_column, 20) + 1):
                    cell = ws.cell(row=r, column=c)
                    cell_data = ws_data.cell(row=r, column=c)
                    val = cell.value
                    val_data = cell_data.value
                    if val is not None or val_data is not None:
                        has_data = True
                        row_data.append({
                            "coord": cell.coordinate,
                            "value": str(val_data),
                            "formula": str(val) if str(val).startswith('=') else None
                        })
                    else:
                        row_data.append(None) # keep alignment
                if has_data:
                    sheet_data.append(row_data)
            out[sheet_name] = sheet_data
        
        with open('excel_formulas.json', 'w', encoding='utf-8') as f:
            json.dump(out, f, ensure_ascii=False, indent=2)
        print("Exported excel_formulas.json")
    except Exception as e:
        print(f"Error: {e}")

if __name__ == '__main__':
    main()
