"""
SRC柱設計 Word 文件產生器
使用 python-docx 產生專業排版的設計計算書
"""
from io import BytesIO

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT


def generate_docx(result):
    """
    根據計算結果產生 Word 文件

    Parameters:
    -----------
    result : dict - calculate_src_column() 的回傳值

    Returns:
    --------
    BytesIO - Word 文件的 bytes buffer
    """
    doc = Document()

    # ===== 頁面設定 =====
    section = doc.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)

    # ===== 樣式設定 =====
    style_normal = doc.styles['Normal']
    style_normal.font.name = '標楷體'
    style_normal.font.size = Pt(12)
    style_normal.paragraph_format.line_spacing = 1.5
    style_normal.paragraph_format.space_after = Pt(2)

    # ===== 標題 =====
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run('SRC柱設計計算書')
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.name = '標楷體'

    # 分隔線
    doc.add_paragraph('─' * 50)

    # ===== 輸入參數摘要 =====
    inp = result['input']
    params_title = doc.add_paragraph()
    run = params_title.add_run('【設計參數】')
    run.font.bold = True
    run.font.size = Pt(14)
    run.font.name = '標楷體'

    param_lines = [
        f"SRC柱尺寸：{inp['col_B']}cm × {inp['col_H']}cm",
        f"柱長度：{inp['col_length']}m",
        f"鋼骨尺寸：□{inp['steel_b']}×{inp['steel_h']}×{inp['steel_t']}mm",
        f"鋼筋配置：{inp['rebar_config']}{inp['rebar_size']}",
        f"有效長度係數：Kx={inp['kx']}　Ky={inp['ky']}",
        f"鋼骨降伏強度Fys={inp['fys']} kgf/cm²　({inp['steel_grade']})",
        f"鋼筋降伏強度Fyr={inp['fyr']} kgf/cm²",
        f"混凝土抗壓強度fc'={inp['fc']} kgf/cm²",
        f"設計軸力Pu={inp['pu']} tf",
        f"設計彎矩Mux={inp['mux']} tf-m　Muy={inp['muy']} tf-m",
    ]
    for line in param_lines:
        p = doc.add_paragraph()
        run = p.add_run(line)
        run.font.name = '標楷體'
        run.font.size = Pt(11)

    doc.add_paragraph('')

    # ===== 計算過程（各章節）=====
    sections = result['lines']
    for section_data in sections:
        # 章節標題
        heading_p = doc.add_paragraph()
        heading_p.paragraph_format.space_before = Pt(12)
        heading_p.paragraph_format.space_after = Pt(6)
        run = heading_p.add_run(section_data['title'])
        run.font.bold = True
        run.font.size = Pt(14)
        run.font.name = '標楷體'

        # 章節內容行
        for line in section_data['lines']:
            if line == '':
                doc.add_paragraph('')
                continue

            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(1)

            # 檢查是否包含 OK/NG 關鍵字，用顏色標記
            if 'NG' in line and ('NG' == line.split('　')[-1] or line.endswith('NG')):
                # 分離 NG 前的文字和 NG
                text_before = line.rsplit('NG', 1)[0]
                run = p.add_run(text_before)
                run.font.name = '標楷體'
                run.font.size = Pt(11)

                run_ng = p.add_run('NG')
                run_ng.font.name = '標楷體'
                run_ng.font.size = Pt(11)
                run_ng.font.bold = True
                run_ng.font.color.rgb = RGBColor(220, 38, 38)  # 紅色
            elif 'OK' in line and ('OK' == line.split('　')[-1] or line.endswith('OK')):
                text_before = line.rsplit('OK', 1)[0]
                run = p.add_run(text_before)
                run.font.name = '標楷體'
                run.font.size = Pt(11)

                run_ok = p.add_run('OK')
                run_ok.font.name = '標楷體'
                run_ok.font.size = Pt(11)
                run_ok.font.bold = True
                run_ok.font.color.rgb = RGBColor(22, 163, 74)  # 綠色
            else:
                run = p.add_run(line)
                run.font.name = '標楷體'
                run.font.size = Pt(11)

    # ===== 頁尾 =====
    doc.add_paragraph('')
    footer_p = doc.add_paragraph()
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer_p.add_run('─ 本文件由 SRC柱設計系統自動產生 ─')
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(150, 150, 150)
    run.font.name = '標楷體'

    # 匯出
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
