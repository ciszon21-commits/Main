from __future__ import annotations

from numbers import Real
from pathlib import Path

from django.conf import settings
from django.views.generic import TemplateView
from openpyxl import load_workbook


def _format_cell(value, *, is_percent: bool = False):
    if value is None:
        return ''
    if isinstance(value, Real) and is_percent:
        return f'{float(value) * 100:.2f}%'
    if isinstance(value, Real) and float(value).is_integer():
        return str(int(value))
    if isinstance(value, Real):
        return f'{float(value):.2f}'
    return str(value)


def _is_red_font(cell) -> bool:
    color = cell.font.color if cell.font else None
    if color is None:
        return False
    if color.type == 'rgb' and color.rgb:
        return color.rgb.upper().endswith('FF0000')
    if color.type == 'indexed' and color.indexed in (2, 10):
        return True
    return False


def _build_sheet_payload(
    ws_formula,
    ws_value,
    sheet_key: str,
    title: str,
    hidden_display_cols: set[int] | None = None,
):
    hidden_display_cols = hidden_display_cols or set()
    max_row = 0
    max_col = 0

    # Find used range based on either formula source or cached value source.
    for r in range(1, ws_formula.max_row + 1):
        for c in range(1, ws_formula.max_column + 1):
            vf = ws_formula.cell(r, c).value
            vv = ws_value.cell(r, c).value
            if vf is not None or vv is not None:
                max_row = max(max_row, r)
                max_col = max(max_col, c)

    headers = [f'欄位 {i + 1}' for i in range(max_col) if (i + 1) not in hidden_display_cols]
    rows = []
    formulas = {}
    cell_values = {}
    input_cells = []
    percent_cells = []

    for r in range(1, max_row + 1):
        row_cells = []
        for c in range(1, max_col + 1):
            cell_f = ws_formula.cell(r, c)
            cell_v = ws_value.cell(r, c)
            addr = cell_f.coordinate

            formula = cell_f.value if isinstance(cell_f.value, str) and cell_f.value.startswith('=') else ''
            # For formula cells, use cached computed value as initial display.
            raw_value = cell_v.value if formula else cell_f.value
            is_input = _is_red_font(cell_f) and not formula and isinstance(raw_value, (int, float))
            number_format = str(cell_f.number_format or '')
            is_percent = '%' in number_format

            if formula:
                formulas[addr] = formula
            cell_values[addr] = raw_value
            if is_input:
                input_cells.append(addr)
            if is_percent:
                percent_cells.append(addr)

            if c not in hidden_display_cols:
                row_cells.append(
                    {
                        'address': addr,
                        'display': _format_cell(raw_value, is_percent=is_percent),
                        'is_input': is_input,
                        'is_formula': bool(formula),
                        'is_percent': is_percent,
                    }
                )

        rows.append(row_cells)

    return {
        'key': sheet_key,
        'title': title,
        'headers': headers,
        'rows': rows,
        'engine': {
            'key': sheet_key,
            'title': title,
            'row_count': max_row,
            'col_count': max_col,
            'formulas': formulas,
            'cell_values': cell_values,
            'input_cells': input_cells,
            'percent_cells': percent_cells,
        },
    }


def _load_building_mass_data(excel_path: Path):
    wb_formula = load_workbook(excel_path, data_only=False)
    wb_value = load_workbook(excel_path, data_only=True)

    sheet_defs = [
        (2, 'taipei', '北市建築量體', {8, 9}),
        (3, 'new_taipei', '新北建築量體', {8, 9, 10}),
    ]

    sheets = []
    for idx, key, title, hidden_cols in sheet_defs:
        sheets.append(
            _build_sheet_payload(
                wb_formula.worksheets[idx],
                wb_value.worksheets[idx],
                key,
                title,
                hidden_display_cols=hidden_cols,
            )
        )

    external_values = {}
    for ws in wb_value.worksheets:
        for r in range(1, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                val = ws.cell(r, c).value
                if val is not None:
                    external_values[f'{ws.title}!{ws.cell(r, c).coordinate}'] = val

    engine_payload = {
        'sheets': {sheet['key']: sheet['engine'] for sheet in sheets},
        'external_values': external_values,
    }

    return sheets, engine_payload


class CalculatorView(TemplateView):
    template_name = 'volume_calc/calculator.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['default_base_area'] = 1029
        context['default_original_volume'] = 2000
        context['default_zone_type'] = '第三種住宅區'
        context['default_volume_ratio'] = 2.25

        context['building_mass_sheets'] = []
        context['building_mass_error'] = ''
        context['building_mass_engine'] = {}

        try:
            reference_dir = Path(settings.BASE_DIR) / 'reference'
            excel_files = sorted(reference_dir.glob('*.xlsx'))
            if not excel_files:
                raise FileNotFoundError('No Excel file found in reference folder.')

            sheets, engine_payload = _load_building_mass_data(excel_files[0])
            context['building_mass_sheets'] = sheets
            context['building_mass_engine'] = engine_payload
        except Exception as exc:
            context['building_mass_error'] = f'讀取建築量體工作表失敗：{exc}'

        return context
