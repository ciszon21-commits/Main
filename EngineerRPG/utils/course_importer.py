"""
課程批次匯入工具
配合 Course 模型結構
"""

import csv
import io
from openpyxl import load_workbook
from django.db import transaction
from EngineerRPG.models import Course

class CourseImporter:
    """課程匯入器"""
    
    def __init__(self):
        print("CourseImporter initialized")
        self.errors = []
        self.success_count = 0
        self.created_count = 0
        self.skip_count = 0
    
    def import_from_file(self, file):
        """從檔案匯入課程"""
        if file.name.endswith('.csv'):
            return self._import_csv(file)
        elif file.name.endswith('.xlsx'):
            return self._import_excel(file)
        else:
            raise ValueError('不支援的檔案格式')
    
    def _import_csv(self, file):
        """匯入 CSV 檔案"""
        decoded_file = file.read().decode('utf-8-sig')
        reader = csv.DictReader(io.StringIO(decoded_file))
        return self._process_rows(reader)
    
    def _import_excel(self, file):
        """匯入 Excel 檔案"""
        wb = load_workbook(file)
        ws = wb.active
        
        # Robust header reading: strip whitespace and ignore None
        headers = [str(cell.value).strip() if cell.value else f"Column_{i}" for i, cell in enumerate(ws[1])]
        
        rows = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            row_dict = dict(zip(headers, row))
            rows.append(row_dict)
        
        return self._process_rows(rows)

    @transaction.atomic
    def _process_rows(self, rows):
        """處理資料列"""
        print(f"Processing {len(rows) if isinstance(rows, list) else 'CSV iterator'} rows")
        for idx, row in enumerate(rows, start=2):
            try:
                # Normalize row keys (support aliases)
                row = self._normalize_row(row)
                
                _, created = self._upsert_course(row)
                if created:
                    self.created_count += 1
                self.success_count += 1
            except Exception as e:
                self.errors.append(f'第 {idx} 列: {str(e)}')
                self.skip_count += 1
        
        return {
            'success': self.success_count,
            'created_count': self.created_count,
            'skip': self.skip_count,
            'errors': self.errors
        }

    def _normalize_row(self, row):
        """標準化欄位名稱 (支援別名)"""
        alias_map = {
            '標題': '課程標題',
            '名稱': '課程標題',
            '描述': '課程描述',
            '類型': '內容類型',
            '連結': '內容網址',
            '網址': '內容網址',
            '時長': '課程時長',
            '時限': '考試時限',
            '分數': '及格分數',
        }
        
        new_row = {}
        for key, value in row.items():
            if not key: continue
            key = str(key).strip()
            new_key = alias_map.get(key, key)
            new_row[new_key] = value
        return new_row
    
    def _upsert_course(self, row):
        """新增或更新課程"""
        required_fields = ['課程標題', '內容類型']
        
        missing = []
        for field in required_fields:
            if not row.get(field):
                missing.append(field)
        
        if missing:
            available_keys = list(row.keys())
            raise ValueError(f'缺少必填欄位: {", ".join(missing)}。 (讀取到的欄位: {available_keys})')
            
        # 準備欄位資料
        defaults = {
            'title': row['課程標題'],
            'description': row.get('課程描述', ''),
            'content_type': self._get_content_type(row['內容類型']),
            'content_url': row.get('內容網址', ''),
            'duration_minutes': int(row.get('課程時長', 30) or 30),
            'exam_time_limit': int(row.get('考試時限', 20) or 20),
            'passing_score': int(row.get('及格分數', 80) or 80),
        }
        
        # 檢查是否有 ID 進行更新
        course_id = row.get('ID')
        if course_id:
            try:
                course = Course.objects.get(id=course_id)
                for key, value in defaults.items():
                    setattr(course, key, value)
                course.save()
                return course, False
            except Course.DoesNotExist:
                pass

        # 建立新課程
        course = Course(**defaults)
        course.save()
        
        return course, True
    
    def _get_content_type(self, type_str):
        """取得內容類型代碼"""
        type_upper = str(type_str).upper().strip()
        if type_upper in ['PDF', 'PDF文件']:
            return 'PDF'
        elif type_upper in ['LINK', '連結', '外部連結']:
            return 'LINK'
        return 'LINK' # Default


def generate_template_excel():
    """產生匯入範本 Excel"""
    from openpyxl import Workbook
    from io import BytesIO
    
    wb = Workbook()
    ws = wb.active
    ws.title = "課程匯入範本"
    
    headers = [
        'ID', '課程標題', '課程描述', '內容類型', 
        '內容網址', '課程時長', '考試時限'
    ]
    ws.append(headers)
    
    examples = [
        ['', '新進人員訓練', '基礎入職培訓課程', 'PDF', '', 60, 30],
        ['', 'Python基礎', 'Python程式設計入門', '連結', 'https://python.org', 120, 40],
    ]
    
    for example in examples:
        ws.append(example)
    
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    return output

