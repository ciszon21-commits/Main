from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Iterable
import xml.etree.ElementTree as ET

from .exceptions import DependencyError, ParseError
from .normalization import parse_decimal, normalize_text
from .schemas import StandardRecord

NAME_ALIASES = {
    'name',
    'item',
    'itemname',
    'description',
    'workitem',
    'itemdescription',
    'itemdesc',
    'workdescription',
    '項目及說明',
    '工作項目',
    '工作項目及說明',
    '項目名稱',
    '項目',
    '工項',
    '名稱',
    '項次名稱',
    '品名',
}
QUANTITY_ALIASES = {
    'qty',
    'quantity',
    'amount',
    '數量',
    '數值',
    '工程數量',
}
UNIT_ALIASES = {
    'unit',
    'uom',
    'measurementunit',
    '計價單位',
    '單位',
}
ROW_ID_ALIASES = {
    'id',
    'no',
    'itemno',
    '項次',
    '編號',
}


@dataclass(frozen=True)
class HeaderMapping:
    name_index: int
    quantity_index: int | None
    unit_index: int | None


def parse_budget_xml(file_name: str, payload: bytes) -> list[StandardRecord]:
    try:
        root = ET.fromstring(payload)
    except ET.ParseError as exc:
        raise ParseError(f'無法解析 XML 檔案: {file_name}') from exc

    results: list[StandardRecord] = []
    auto_id = 1

    for element in root.iter():
        candidate = _build_candidate_from_xml_element(element)
        name = _first_non_empty(candidate, NAME_ALIASES)
        quantity = _first_non_empty(candidate, QUANTITY_ALIASES)
        unit = _first_non_empty(candidate, UNIT_ALIASES)
        if not name:
            continue
        parsed_qty = parse_decimal(quantity)
        if parsed_qty is None and not unit:
            continue

        row_id = _first_non_empty(candidate, ROW_ID_ALIASES) or f'xml-{auto_id}'
        auto_id += 1
        results.append(
            StandardRecord(
                source_type='budget_xml',
                source_id=str(row_id),
                name=name,
                quantity=parsed_qty,
                unit=unit or '',
                raw_payload=candidate,
            )
        )

    if not results:
        raise ParseError(
            'XML 內找不到可辨識的工項資料。請確認節點包含名稱/數量/單位欄位。'
        )
    return results


def parse_quantity_workbook(file_name: str, payload: bytes) -> list[StandardRecord]:
    suffix = Path(file_name).suffix.lower()
    if suffix == '.xlsx':
        return _parse_xlsx(file_name, payload)
    if suffix == '.xls':
        return _parse_xls(file_name, payload)
    raise ParseError(f'不支援的數量計算書格式: {file_name}')


def _parse_xlsx(file_name: str, payload: bytes) -> list[StandardRecord]:
    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError as exc:
        raise DependencyError('缺少套件 openpyxl，請先安裝: pip install openpyxl') from exc

    workbook = load_workbook(filename=BytesIO(payload), read_only=True, data_only=True)
    results: list[StandardRecord] = []

    for worksheet in workbook.worksheets:
        rows = list(worksheet.iter_rows(values_only=True))
        sheet_results = _parse_sheet_rows(rows, worksheet.title)
        results.extend(sheet_results)

    if not results:
        raise ParseError(f'Excel 檔案內找不到可辨識的工項資料: {file_name}')
    return results


def _parse_xls(file_name: str, payload: bytes) -> list[StandardRecord]:
    try:
        import xlrd
    except ModuleNotFoundError as exc:
        raise DependencyError('缺少套件 xlrd，請先安裝: pip install xlrd') from exc

    workbook = xlrd.open_workbook(file_contents=payload)
    results: list[StandardRecord] = []

    for sheet in workbook.sheets():
        rows = [sheet.row_values(i) for i in range(sheet.nrows)]
        sheet_results = _parse_sheet_rows(rows, sheet.name)
        results.extend(sheet_results)

    if not results:
        raise ParseError(f'Excel 檔案內找不到可辨識的工項資料: {file_name}')
    return results


def _parse_sheet_rows(rows: list[Iterable[object]], sheet_name: str) -> list[StandardRecord]:
    if not rows:
        return []

    header_index, mapping = _detect_header(rows)
    if mapping is None:
        return []

    records: list[StandardRecord] = []
    for row_number, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        row_values = list(row)
        raw_name = _safe_get(row_values, mapping.name_index)
        if raw_name is None or str(raw_name).strip() == '':
            continue

        raw_qty = _safe_get(row_values, mapping.quantity_index) if mapping.quantity_index is not None else None
        raw_unit = _safe_get(row_values, mapping.unit_index) if mapping.unit_index is not None else ''
        quantity = parse_decimal(raw_qty)
        raw_unit_text = '' if raw_unit is None else str(raw_unit).strip()

        # Skip heading rows without quantity and unit.
        if quantity is None and not raw_unit_text:
            continue

        records.append(
            StandardRecord(
                source_type='quantity_sheet',
                source_id=f'{sheet_name}:{row_number}',
                name=str(raw_name).strip(),
                quantity=quantity,
                unit=raw_unit_text,
                raw_payload={
                    'sheet': sheet_name,
                    'row_number': row_number,
                    'raw_name': raw_name,
                    'raw_qty': raw_qty,
                    'raw_unit': raw_unit,
                },
            )
        )

    return records


def _detect_header(rows: list[Iterable[object]]) -> tuple[int, HeaderMapping | None]:
    best_index = -1
    best_mapping: HeaderMapping | None = None
    best_score = -1

    for idx, row in enumerate(rows[:120]):
        values = [normalize_text(str(cell)) for cell in list(row)]
        candidate = _build_header_mapping(values)
        if candidate is None:
            continue
        score = 1
        if candidate.quantity_index is not None:
            score += 1
        if candidate.unit_index is not None:
            score += 1

        if score > best_score:
            best_score = score
            best_index = idx
            best_mapping = candidate

    return best_index, best_mapping


def _build_header_mapping(values: list[str]) -> HeaderMapping | None:
    name_idx = _find_index(values, NAME_ALIASES)
    qty_idx = _find_index(values, QUANTITY_ALIASES)
    unit_idx = _find_index(values, UNIT_ALIASES)

    if name_idx is None:
        return None
    if qty_idx is None and unit_idx is None:
        return None

    return HeaderMapping(name_index=name_idx, quantity_index=qty_idx, unit_index=unit_idx)


def _build_candidate_from_xml_element(element: ET.Element) -> dict[str, str]:
    data: dict[str, str] = {}
    selected_language: dict[str, str] = {}
    for key, value in element.attrib.items():
        clean_key = _strip_ns(key)
        data[clean_key] = str(value).strip()

    for child in list(element):
        if list(child):
            continue
        if child.text is None:
            continue
        child_text = child.text.strip()
        if not child_text:
            continue
        child_key = _strip_ns(child.tag)
        language = normalize_text(str(child.attrib.get('language', ''))).replace('_', '-')
        if language:
            data[f'{child_key}[{language}]'] = child_text
            previous_language = selected_language.get(child_key, '')
            if _language_rank(language) >= _language_rank(previous_language):
                data[child_key] = child_text
                selected_language[child_key] = language
        elif child_key not in data:
            data[child_key] = child_text

    # Keep an optional value of element text for very flat XML structures.
    if element.text and element.text.strip():
        data[_strip_ns(element.tag)] = element.text.strip()

    return data


def _strip_ns(tag: str) -> str:
    if '}' in tag:
        return tag.split('}', 1)[1]
    return tag


def _language_rank(language: str) -> int:
    normalized = normalize_text(language).replace('_', '-')
    if normalized in {'zh-tw', 'zh-hant'}:
        return 30
    if normalized.startswith('zh'):
        return 20
    if normalized:
        return 10
    return 0


def _find_index(values: list[str], aliases: set[str]) -> int | None:
    normalized_aliases = {_normalize_header_text(alias) for alias in aliases if alias}

    for idx, value in enumerate(values):
        compact = _normalize_header_text(value)
        if compact in normalized_aliases:
            return idx

    # Fallback: allow partial match (e.g. "項目及說明" contains "項目").
    candidate_aliases = [alias for alias in normalized_aliases if len(alias) >= 2]
    for idx, value in enumerate(values):
        compact = _normalize_header_text(value)
        if not compact:
            continue
        if any(alias in compact for alias in candidate_aliases):
            return idx
    return None


def _normalize_header_text(value: str) -> str:
    compact = normalize_text(value)
    for char in (' ', '\n', '\r', '\t', ':', '：', '(', ')', '[', ']', '【', '】'):
        compact = compact.replace(char, '')
    return compact


def _safe_get(values: list[object], index: int | None) -> object | None:
    if index is None:
        return None
    if index < 0 or index >= len(values):
        return None
    return values[index]


def _first_non_empty(data: dict[str, str], aliases: set[str]) -> str:
    for key, value in data.items():
        normalized_key = normalize_text(key).replace(' ', '')
        if normalized_key in aliases and str(value).strip():
            return str(value).strip()
    return ''
