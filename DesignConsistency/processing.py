import difflib
import re
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any
from datetime import datetime

from django.conf import settings

import openpyxl
from openpyxl import Workbook

try:
    from rapidfuzz import fuzz as rf_fuzz
    from rapidfuzz import process as rf_process
except ImportError:  # pragma: no cover
    rf_fuzz = None
    rf_process = None

try:
    from lxml import etree
except ImportError:  # pragma: no cover
    etree = None


ETENDER_NS = {"ns": "http://pcstd.pcc.gov.tw/2003/eTender"}


HEADER_KEYWORDS = {
    "item_no": ["項次", "項號", "序號", "編號", "項目編號", "項目代碼"],
    "description": ["項目", "名稱", "說明", "內容", "品名", "工程項目", "工作項目"],
    "unit": ["單位", "unit", "單　位"],
    "quantity": ["數量", "數　量", "qty", "數 量"],
    "category": ["分類", "類別"],
}

MAX_CANDIDATES = 120
MIN_TOKEN_OVERLAP = 0.45
MIN_COMMON_TOKENS = 2
HIGH_FREQ_RATIO = 0.35

FULLWIDTH_DIGITS = str.maketrans(
    {
        "\uff10": "0",
        "\uff11": "1",
        "\uff12": "2",
        "\uff13": "3",
        "\uff14": "4",
        "\uff15": "5",
        "\uff16": "6",
        "\uff17": "7",
        "\uff18": "8",
        "\uff19": "9",
    }
)

NUMBER_TOKEN_PATTERN = re.compile(
    r"\d+(?:\.\d+)?(?:\s*(?:mm|cm|m|kg|kn|ton|t|m2|m3|%))?",
    re.IGNORECASE,
)


def _normalize_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    text = text.replace("\u3000", " ")
    text = re.sub(r"\s+", " ", text)
    return text


def _normalize_key(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    text = text.replace("\u3000", "")
    text = re.sub(r"\s+", "", text)
    return text


def _parse_number(value: Any) -> float | None:
    if value is None:
        return None
    text = _normalize_key(value)
    if not text:
        return None
    try:
        return float(text.replace(",", ""))
    except Exception:
        return None


def _tokenize(text: str) -> list[str]:
    if not text:
        return []
    tokens = []
    tokens.extend(re.findall(r"[\u4e00-\u9fff]{2,}", text))
    tokens.extend(re.findall(r"[A-Za-z][A-Za-z0-9_/\\-]{1,}", text))
    normalized = []
    for token in tokens:
        if token.isascii():
            token = token.lower()
        if len(token) < 2:
            continue
        if token.isdigit():
            continue
        normalized.append(token)
    return normalized


def _apply_synonyms(tokens: list[str], synonym_map: dict[str, str] | None) -> list[str]:
    if not synonym_map:
        return tokens
    return [synonym_map.get(token, token) for token in tokens]


def normalize_for_match(text: str, synonym_map: dict[str, str] | None = None) -> str:
    tokens = _tokenize(text)
    tokens = _apply_synonyms(tokens, synonym_map)
    return " ".join(tokens)


def _normalize_fullwidth_digits(text: str) -> str:
    return text.translate(FULLWIDTH_DIGITS)


def _extract_number_tokens(text: str) -> list[str]:
    if not text:
        return []
    text = _normalize_fullwidth_digits(text)
    tokens: list[str] = []
    for match in NUMBER_TOKEN_PATTERN.finditer(text):
        token = re.sub(r"\s+", "", match.group(0))
        if token:
            tokens.append(token.lower())
    for match in re.finditer(r"\d+(?:\.\d+)?", text):
        token = match.group(0)
        if token:
            tokens.append(token)
    seen: set[str] = set()
    result: list[str] = []
    for token in tokens:
        if token in seen:
            continue
        seen.add(token)
        result.append(token)
    return result


def _build_text_profile(text: str, synonym_map: dict[str, str] | None) -> dict[str, Any]:
    raw = _normalize_text(text)
    tokens = _apply_synonyms(_tokenize(raw), synonym_map)
    token_set = set(tokens)
    signature = " ".join(sorted(token_set))
    numbers = _extract_number_tokens(raw)
    return {
        "raw": raw,
        "tokens": tokens,
        "token_set": token_set,
        "signature": signature,
        "norm": " ".join(tokens),
        "numbers": numbers,
        "number_set": set(numbers),
    }


def _score_profiles(q_profile: dict[str, Any], b_profile: dict[str, Any]) -> tuple[int, int, float, set[str]]:
    if rf_fuzz:
        base_score = int(rf_fuzz.token_set_ratio(q_profile["norm"], b_profile["norm"]))
    else:
        base_score = int(
            round(
                100
                * difflib.SequenceMatcher(None, q_profile["norm"], b_profile["norm"]).ratio()
            )
        )
    q_tokens = q_profile["token_set"]
    b_tokens = b_profile["token_set"]
    common = q_tokens & b_tokens
    overlap = len(common) / max(1, min(len(q_tokens), len(b_tokens)))
    score = int(round(base_score * 0.7 + overlap * 100 * 0.3))
    return score, base_score, overlap, common


def _build_name_diff(
    q_profile: dict[str, Any], b_profile: dict[str, Any], limit: int = 12
) -> dict[str, Any]:
    tokens_only_budget = sorted(b_profile["token_set"] - q_profile["token_set"])
    tokens_only_quantity = sorted(q_profile["token_set"] - b_profile["token_set"])
    numbers_only_budget = sorted(b_profile["number_set"] - q_profile["number_set"])
    numbers_only_quantity = sorted(q_profile["number_set"] - b_profile["number_set"])
    return {
        "tokens_only_budget": tokens_only_budget[:limit],
        "tokens_only_quantity": tokens_only_quantity[:limit],
        "numbers_only_budget": numbers_only_budget[:limit],
        "numbers_only_quantity": numbers_only_quantity[:limit],
        "has_diff": bool(
            tokens_only_budget
            or tokens_only_quantity
            or numbers_only_budget
            or numbers_only_quantity
        ),
    }


def _summarize_list(values: list[str], limit: int = 8) -> str:
    if not values:
        return ""
    trimmed = values[:limit]
    suffix = "..." if len(values) > limit else ""
    return ", ".join(trimmed) + suffix


def build_term_stats(texts: list[str], synonym_map: dict[str, str] | None = None) -> dict:
    counter = Counter()
    for text in texts:
        tokens = _apply_synonyms(_tokenize(text), synonym_map)
        counter.update(tokens)
    return {
        "total_terms": len(counter),
        "top_terms": [{"term": term, "count": count} for term, count in counter.most_common(20)],
    }


def build_budget_stats(budget_rows: list[dict[str, Any]]) -> dict[str, Any]:
    kind_counts = Counter()
    kind_unique = defaultdict(set)
    unique_desc = set()
    qty_only = 0
    analysis_only = 0
    for row in budget_rows:
        kind = (row.get("item_kind") or "unknown").lower()
        kind_counts[kind] += 1
        desc_key = _normalize_key(row.get("description"))
        if desc_key:
            unique_desc.add(desc_key)
            kind_unique[kind].add(desc_key)
        if row.get("quantity") is not None:
            if kind not in {"mainitem", "subtotal"}:
                qty_only += 1
            if kind in {"analysis", "general"}:
                analysis_only += 1
    total = len(budget_rows)
    kind_labels = {
        "mainitem": "主項",
        "analysis": "分析",
        "subtotal": "小計",
        "general": "一般",
        "unknown": "未知",
    }
    by_kind = dict(kind_counts)
    by_kind_cn = []
    for kind, count in sorted(by_kind.items(), key=lambda item: (-item[1], item[0])):
        unique_count = len(kind_unique.get(kind, set()))
        unique_ratio = (unique_count / count) if count else 0
        by_kind_cn.append(
            {
                "kind": kind,
                "label": kind_labels.get(kind, kind),
                "count": count,
                "unique_count": unique_count,
                "unique_ratio": round(unique_ratio * 100, 2),
            }
        )
    qty_only_ratio = (qty_only / total) if total else 0
    analysis_only_ratio = (analysis_only / total) if total else 0
    unique_ratio = (len(unique_desc) / total) if total else 0
    return {
        "total": total,
        "unique_total": len(unique_desc),
        "by_kind": by_kind,
        "by_kind_cn": by_kind_cn,
        "quantity_only": qty_only,
        "analysis_only": analysis_only,
        "quantity_only_ratio": round(qty_only_ratio * 100, 2),
        "analysis_only_ratio": round(analysis_only_ratio * 100, 2),
        "unique_ratio": round(unique_ratio * 100, 2),
    }


def build_quantity_stats(quantity_rows: list[dict[str, Any]]) -> dict[str, Any]:
    sheet_counts = Counter()
    sheet_unique = defaultdict(set)
    unique_desc = set()
    for row in quantity_rows:
        sheet = row.get("sheet") or "unknown"
        sheet_counts[sheet] += 1
        desc_key = _normalize_key(row.get("description"))
        if desc_key:
            unique_desc.add(desc_key)
            sheet_unique[sheet].add(desc_key)
    total = len(quantity_rows)
    by_sheet_cn = []
    for sheet, count in sorted(sheet_counts.items(), key=lambda item: (-item[1], item[0])):
        unique_count = len(sheet_unique.get(sheet, set()))
        unique_ratio = (unique_count / count) if count else 0
        by_sheet_cn.append(
            {
                "sheet": sheet,
                "count": count,
                "unique_count": unique_count,
                "unique_ratio": round(unique_ratio * 100, 2),
            }
        )
    unique_ratio = (len(unique_desc) / total) if total else 0
    return {
        "total": total,
        "unique_total": len(unique_desc),
        "unique_ratio": round(unique_ratio * 100, 2),
        "by_sheet": dict(sheet_counts),
        "by_sheet_cn": by_sheet_cn,
    }


def build_compare_suggestion(
    budget_stats: dict[str, Any] | None,
    quantity_stats: dict[str, Any] | None,
) -> dict[str, str]:
    if not budget_stats or not quantity_stats:
        return {"mode": "balanced", "budget_filter": "", "reason": "資料不足，使用預設模式。"}

    budget_total = budget_stats.get("total") or 0
    budget_qty = budget_stats.get("quantity_only") or 0
    quantity_total = quantity_stats.get("total") or 0

    if budget_total <= 0 or quantity_total <= 0:
        return {"mode": "balanced", "budget_filter": "", "reason": "資料不足，使用預設模式。"}

    ratio = budget_qty / max(1, quantity_total)
    filter_mode = ""
    if budget_total > 0 and (budget_total - budget_qty) / budget_total > 0.2:
        filter_mode = "quantity_only"

    if ratio >= 30:
        mode = "strict"
        reason = "預算可比對項目遠多於數量項目，建議採嚴格匹配降低誤配。"
    elif ratio >= 10:
        mode = "balanced"
        reason = "預算項目明顯多於數量項目，使用平衡模式較穩定。"
    elif ratio <= 1.5:
        mode = "lenient"
        reason = "數量項目與預算項目接近，建議使用寬鬆模式提高召回。"
    else:
        mode = "balanced"
        reason = "項目比例中等，使用平衡模式。"

    return {"mode": mode, "budget_filter": filter_mode, "reason": reason}


def build_stats_workbook(
    project_name: str,
    budget_stats: dict[str, Any] | None,
    quantity_stats: dict[str, Any] | None,
) -> Workbook:
    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = "Summary"
    ws_summary.append(["專案", project_name])

    if budget_stats:
        ws_summary.append([])
        ws_summary.append(["預算書統計"])
        ws_summary.append(["總項目數", budget_stats.get("total")])
        ws_summary.append(["名稱去重", budget_stats.get("unique_total")])
        ws_summary.append(["名稱去重佔比", budget_stats.get("unique_ratio")])
        ws_summary.append(["只比對有數量", budget_stats.get("quantity_only")])
        ws_summary.append(["分析/一般項且有數量", budget_stats.get("analysis_only")])

    if quantity_stats:
        ws_summary.append([])
        ws_summary.append(["數量計算書統計"])
        ws_summary.append(["總項目數", quantity_stats.get("total")])
        ws_summary.append(["名稱去重", quantity_stats.get("unique_total")])
        ws_summary.append(["名稱去重佔比", quantity_stats.get("unique_ratio")])

    ws_budget = wb.create_sheet("BudgetKinds")
    ws_budget.append(["類型", "項目數", "名稱去重", "去重佔比"])
    if budget_stats:
        for item in budget_stats.get("by_kind_cn", []):
            ws_budget.append(
                [
                    item.get("label"),
                    item.get("count"),
                    item.get("unique_count"),
                    item.get("unique_ratio"),
                ]
            )

    ws_qty = wb.create_sheet("QuantitySheets")
    ws_qty.append(["工作表", "項目數", "名稱去重", "去重佔比"])
    if quantity_stats:
        for item in quantity_stats.get("by_sheet_cn", []):
            ws_qty.append(
                [
                    item.get("sheet"),
                    item.get("count"),
                    item.get("unique_count"),
                    item.get("unique_ratio"),
                ]
            )
    return wb


def _detect_header_row(ws, max_scan: int = 20) -> tuple[int, dict[str, int]]:
    best_row = 1
    best_score = -1
    best_map: dict[str, int] = {}
    max_row = min(ws.max_row, max_scan)
    max_col = ws.max_column

    for row_idx in range(1, max_row + 1):
        row_values = [
            _normalize_key(ws.cell(row=row_idx, column=col).value)
            for col in range(1, max_col + 1)
        ]
        mapping: dict[str, int] = {}
        for key, keywords in HEADER_KEYWORDS.items():
            for col_idx, cell_val in enumerate(row_values, start=1):
                for kw in keywords:
                    if kw and kw in cell_val:
                        mapping[key] = col_idx
                        break
                if key in mapping:
                    break
        score = sum(1 for k in ("description", "unit", "quantity") if k in mapping)
        if score > best_score:
            best_score = score
            best_row = row_idx
            best_map = mapping

    return best_row, best_map


def _detect_header_row_values(values: list[list[Any]], max_scan: int = 20) -> tuple[int, dict[str, int]]:
    best_row = 0
    best_score = -1
    best_map: dict[str, int] = {}
    scan_rows = min(len(values), max_scan)
    for row_idx in range(scan_rows):
        row_values = [_normalize_key(val) for val in values[row_idx]]
        mapping: dict[str, int] = {}
        for key, keywords in HEADER_KEYWORDS.items():
            for col_idx, cell_val in enumerate(row_values, start=1):
                for kw in keywords:
                    if kw and kw in cell_val:
                        mapping[key] = col_idx
                        break
                if key in mapping:
                    break
        score = sum(1 for k in ("description", "unit", "quantity") if k in mapping)
        if score > best_score:
            best_score = score
            best_row = row_idx
            best_map = mapping
    return best_row + 1, best_map


def extract_quantity_rows(file_path: str) -> list[dict[str, Any]]:
    ext = Path(file_path).suffix.lower()
    if ext == ".xls":
        try:
            import pandas as pd
        except ImportError as exc:
            raise RuntimeError("Missing pandas for .xls parsing.") from exc
        sheets = pd.read_excel(file_path, sheet_name=None, header=None, engine="xlrd")
        rows: list[dict[str, Any]] = []
        for sheet_name, frame in sheets.items():
            values = frame.fillna("").values.tolist()
            header_row, header_map = _detect_header_row_values(values)
            if "description" not in header_map or "quantity" not in header_map:
                continue
            current_category = ""
            for idx in range(header_row, len(values)):
                row_vals = values[idx]
                desc = _normalize_text(row_vals[header_map["description"] - 1] if header_map.get("description") else "")
                qty = _parse_number(row_vals[header_map["quantity"] - 1] if header_map.get("quantity") else None)
                unit = _normalize_text(row_vals[header_map["unit"] - 1] if header_map.get("unit") else "")
                item_no = _normalize_text(row_vals[header_map["item_no"] - 1] if header_map.get("item_no") else "")
                category = _normalize_text(row_vals[header_map["category"] - 1] if header_map.get("category") else "")

                if desc and qty is None and category and not item_no:
                    current_category = category
                    continue
                if not desc or qty is None:
                    continue
                rows.append(
                    {
                        "item_no": item_no,
                        "category": category or current_category,
                        "description": desc,
                        "unit": unit,
                        "quantity": qty,
                        "sheet": sheet_name,
                        "row": idx + 1,
                    }
                )
        return rows

    wb = openpyxl.load_workbook(file_path, data_only=True)
    rows: list[dict[str, Any]] = []
    for ws in wb.worksheets:
        header_row, header_map = _detect_header_row(ws)
        if "description" not in header_map or "quantity" not in header_map:
            continue
        current_category = ""
        for r in range(header_row + 1, ws.max_row + 1):
            desc = _normalize_text(ws.cell(row=r, column=header_map["description"]).value)
            qty = _parse_number(ws.cell(row=r, column=header_map["quantity"]).value)
            unit_col = header_map.get("unit")
            item_col = header_map.get("item_no")
            category_col = header_map.get("category")
            unit = _normalize_text(ws.cell(row=r, column=unit_col).value) if unit_col else ""
            item_no = _normalize_text(ws.cell(row=r, column=item_col).value) if item_col else ""
            category = _normalize_text(ws.cell(row=r, column=category_col).value) if category_col else ""

            if desc and qty is None and category and not item_no:
                current_category = category
                continue
            if not desc or qty is None:
                continue
            rows.append(
                {
                    "item_no": item_no,
                    "category": category or current_category,
                    "description": desc,
                    "unit": unit,
                    "quantity": qty,
                    "sheet": ws.title,
                    "row": r,
                }
            )
    return rows


def _get_text(elem, tag: str, lang: str | None = None) -> str:
    if etree is None:
        raise RuntimeError("Missing lxml for XML parsing.")
    if lang:
        node = elem.find(f'ns:{tag}[@language="{lang}"]', ETENDER_NS)
    else:
        node = elem.find(f"ns:{tag}", ETENDER_NS)
    if node is not None and node.text:
        return _normalize_text(node.text)
    return ""


def _build_budget_row(elem) -> dict[str, Any]:
    return {
        "item_no": elem.get("itemNo", "") or elem.get("itemCode", ""),
        "ref_code": elem.get("refItemCode", "") or elem.get("refItemNo", ""),
        "item_kind": (elem.get("itemKind") or "").strip().lower(),
        "description": _get_text(elem, "Description", lang="zh-TW"),
        "unit": _get_text(elem, "Unit", lang="zh-TW"),
        "quantity": _parse_number(_get_text(elem, "Quantity")),
        "price": _parse_number(_get_text(elem, "Price")),
        "amount": _parse_number(_get_text(elem, "Amount")),
    }


def extract_budget_rows(file_path: str) -> list[dict[str, Any]]:
    if etree is None:
        raise RuntimeError("Missing lxml for XML parsing.")
    parser = etree.XMLParser(recover=True, huge_tree=True)
    tree = etree.parse(file_path, parser)
    root = tree.getroot()

    rows: list[dict[str, Any]] = []
    detail_list = root.find(".//ns:DetailList", ETENDER_NS)
    if detail_list is not None:
        for elem in detail_list.findall("ns:PayItem", ETENDER_NS):
            rows.extend(_walk_xml_items(elem))
    else:
        for elem in root.findall(".//ns:PayItem", ETENDER_NS):
            rows.extend(_walk_xml_items(elem))

    cbl = root.find(".//ns:CostBreakdownList", ETENDER_NS)
    if cbl is not None:
        for elem in cbl.findall("ns:WorkItem", ETENDER_NS):
            rows.extend(_walk_xml_items(elem))
    else:
        for elem in root.findall(".//ns:WorkItem", ETENDER_NS):
            rows.extend(_walk_xml_items(elem))

    return rows


def _walk_xml_items(elem) -> list[dict[str, Any]]:
    rows = [_build_budget_row(elem)]
    for child in elem:
        if child.tag.endswith("PayItem") or child.tag.endswith("WorkItem"):
            rows.extend(_walk_xml_items(child))
    return rows


def _filter_budget_rows(
    budget_rows: list[dict[str, Any]],
    budget_filter: str | None,
) -> tuple[list[dict[str, Any]], dict[str, int]]:
    mode = (budget_filter or "").strip().lower()
    if not mode or mode == "none":
        return budget_rows, {"total": len(budget_rows), "kept": len(budget_rows), "skipped": 0}

    kept: list[dict[str, Any]] = []
    for row in budget_rows:
        item_kind = (row.get("item_kind") or "").lower()
        qty = row.get("quantity")
        if mode == "quantity_only":
            if qty is None:
                continue
            if item_kind in {"mainitem", "subtotal"}:
                continue
            kept.append(row)
        elif mode == "analysis_only":
            if item_kind and item_kind not in {"analysis", "general"}:
                continue
            if qty is None:
                continue
            kept.append(row)
        else:
            kept.append(row)

    return kept, {"total": len(budget_rows), "kept": len(kept), "skipped": len(budget_rows) - len(kept)}


def compare_rows_classic(
    budget_rows: list[dict[str, Any]],
    quantity_rows: list[dict[str, Any]],
    synonym_map: dict[str, str] | None = None,
    threshold: int = 88,
) -> dict:
    budget_norms = [normalize_for_match(row.get("description", ""), synonym_map) for row in budget_rows]
    quantity_norms = [normalize_for_match(row.get("description", ""), synonym_map) for row in quantity_rows]

    budget_lookup: dict[str, list[int]] = {}
    for idx, key in enumerate(budget_norms):
        if not key:
            continue
        budget_lookup.setdefault(key, []).append(idx)

    used_budget: set[int] = set()
    matches: list[dict[str, Any]] = []
    unmatched_quantity: list[int] = []

    for q_idx, q_key in enumerate(quantity_norms):
        if not q_key:
            unmatched_quantity.append(q_idx)
            continue
        match_idx = None
        match_score = 0
        match_method = "none"

        if q_key in budget_lookup:
            candidates = budget_lookup[q_key]
            for idx in candidates:
                if idx not in used_budget:
                    match_idx = idx
                    break
            if match_idx is None:
                match_idx = candidates[0]
            match_score = 100
            match_method = "exact"
        elif rf_process and rf_fuzz:
            best = rf_process.extractOne(q_key, budget_norms, scorer=rf_fuzz.token_set_ratio)
            if best:
                match_idx, match_score = best[2], int(best[1])
                if match_score >= threshold:
                    match_method = "fuzzy"
                else:
                    match_idx = None

        if match_idx is None:
            unmatched_quantity.append(q_idx)
            continue

        used_budget.add(match_idx)
        q_profile = _build_text_profile(quantity_rows[q_idx].get("description", ""), synonym_map)
        b_profile = _build_text_profile(budget_rows[match_idx].get("description", ""), synonym_map)
        name_diff = _build_name_diff(q_profile, b_profile)
        matches.append(
            {
                "quantity_index": q_idx,
                "budget_index": match_idx,
                "score": match_score,
                "base_score": match_score,
                "method": match_method,
                "token_overlap": None,
                "common_tokens": [],
                "name_diff": name_diff,
            }
        )

    unmatched_budget = [idx for idx in range(len(budget_rows)) if idx not in used_budget]

    summary = {
        "budget_rows": len(budget_rows),
        "quantity_rows": len(quantity_rows),
        "matched_rows": len(matches),
        "unmatched_budget": len(unmatched_budget),
        "unmatched_quantity": len(unmatched_quantity),
    }
    return {
        "summary": summary,
        "matches": matches,
        "unmatched_budget": unmatched_budget,
        "unmatched_quantity": unmatched_quantity,
    }


def compare_rows_token(
    budget_rows: list[dict[str, Any]],
    quantity_rows: list[dict[str, Any]],
    synonym_map: dict[str, str] | None = None,
    threshold: int = 88,
    allow_fuzzy: bool = True,
    allow_global_fuzzy: bool = True,
    min_overlap: float = MIN_TOKEN_OVERLAP,
    min_common_tokens: int = MIN_COMMON_TOKENS,
    max_candidates: int = MAX_CANDIDATES,
) -> dict:
    budget_profiles: list[dict[str, Any]] = []
    budget_signature_lookup: dict[str, list[int]] = defaultdict(list)
    budget_token_index: dict[str, set[int]] = defaultdict(set)
    budget_item_lookup: dict[str, list[int]] = defaultdict(list)
    budget_norms: list[str] = []

    for idx, row in enumerate(budget_rows):
        profile = _build_text_profile(row.get("description", ""), synonym_map)
        budget_profiles.append(profile)
        budget_norms.append(profile["norm"])
        if profile["signature"]:
            budget_signature_lookup[profile["signature"]].append(idx)
        for token in profile["token_set"]:
            budget_token_index[token].add(idx)
        item_key = _normalize_key(row.get("item_no") or row.get("ref_code"))
        if item_key:
            budget_item_lookup[item_key].append(idx)

    token_counts = Counter()
    for profile in budget_profiles:
        token_counts.update(profile["token_set"])
    high_freq_tokens = {
        token
        for token, count in token_counts.items()
        if budget_profiles and (count / len(budget_profiles)) > HIGH_FREQ_RATIO
    }

    used_budget: set[int] = set()
    matches: list[dict[str, Any]] = []
    unmatched_quantity: list[int] = []

    for q_idx, row in enumerate(quantity_rows):
        q_profile = _build_text_profile(row.get("description", ""), synonym_map)
        if not q_profile["signature"]:
            unmatched_quantity.append(q_idx)
            continue

        match_idx = None
        match_score = 0
        base_score = 0
        token_overlap = 0.0
        common_tokens: set[str] = set()
        match_method = "none"

        q_item_key = _normalize_key(row.get("item_no"))
        if q_item_key and q_item_key in budget_item_lookup:
            candidates = [idx for idx in budget_item_lookup[q_item_key] if idx not in used_budget]
            if not candidates:
                candidates = budget_item_lookup[q_item_key]
            best_idx = None
            best_score = -1
            best_overlap = 0.0
            best_common: set[str] = set()
            for idx in candidates:
                score, base, overlap, common = _score_profiles(q_profile, budget_profiles[idx])
                if score > best_score:
                    best_score = score
                    base_score = base
                    best_overlap = overlap
                    best_common = common
                    best_idx = idx
            if best_idx is not None:
                match_idx = best_idx
                match_score = max(best_score, 95)
                token_overlap = best_overlap
                common_tokens = best_common
                match_method = "item_no"

        if match_idx is None and q_profile["signature"] in budget_signature_lookup:
            candidates = [
                idx
                for idx in budget_signature_lookup[q_profile["signature"]]
                if idx not in used_budget
            ]
            if not candidates:
                candidates = budget_signature_lookup[q_profile["signature"]]
            best_idx = None
            best_score = -1
            best_overlap = 0.0
            best_common: set[str] = set()
            for idx in candidates:
                score, base, overlap, common = _score_profiles(q_profile, budget_profiles[idx])
                if score > best_score:
                    best_score = score
                    base_score = base
                    best_overlap = overlap
                    best_common = common
                    best_idx = idx
            if best_idx is not None:
                match_idx = best_idx
                match_score = max(best_score, 98)
                token_overlap = best_overlap
                common_tokens = best_common
                match_method = "signature"

        if match_idx is None and allow_fuzzy:
            candidate_indices: set[int] = set()
            for token in q_profile["token_set"]:
                if token in high_freq_tokens:
                    continue
                candidate_indices.update(budget_token_index.get(token, set()))

            if candidate_indices:
                scored_candidates: list[tuple[float, int]] = []
                for idx in candidate_indices:
                    common = q_profile["token_set"] & budget_profiles[idx]["token_set"]
                    if not common:
                        continue
                    overlap = len(common) / max(
                        1, min(len(q_profile["token_set"]), len(budget_profiles[idx]["token_set"]))
                    )
                    scored_candidates.append((overlap, idx))
                scored_candidates.sort(reverse=True)
                candidate_indices = {idx for _, idx in scored_candidates[:max_candidates]}

            if candidate_indices:
                best_idx = None
                best_score = -1
                best_overlap = 0.0
                best_common: set[str] = set()
                best_base = 0
                for idx in candidate_indices:
                    score, base, overlap, common = _score_profiles(q_profile, budget_profiles[idx])
                    if score > best_score:
                        best_score = score
                        best_base = base
                        best_overlap = overlap
                        best_common = common
                        best_idx = idx
                if best_idx is not None:
                    match_idx = best_idx
                    match_score = best_score
                    base_score = best_base
                    token_overlap = best_overlap
                    common_tokens = best_common
                    match_method = "fuzzy"

        if match_idx is None and allow_global_fuzzy and rf_process and rf_fuzz:
            best = rf_process.extractOne(q_profile["norm"], budget_norms, scorer=rf_fuzz.token_set_ratio)
            if best:
                match_idx, match_score = best[2], int(best[1])
                base_score = match_score
                _, _, token_overlap, common_tokens = _score_profiles(q_profile, budget_profiles[match_idx])
                match_method = "global_fuzzy"

        if match_idx is None:
            unmatched_quantity.append(q_idx)
            continue

        if match_method in {"fuzzy", "global_fuzzy"}:
            if (
                match_score < threshold
                and (token_overlap < min_overlap or len(common_tokens) < min_common_tokens)
            ):
                unmatched_quantity.append(q_idx)
                continue

        used_budget.add(match_idx)
        name_diff = _build_name_diff(q_profile, budget_profiles[match_idx])
        matches.append(
            {
                "quantity_index": q_idx,
                "budget_index": match_idx,
                "score": match_score,
                "base_score": base_score,
                "method": match_method,
                "token_overlap": round(token_overlap, 3),
                "common_tokens": sorted(list(common_tokens))[:20],
                "name_diff": name_diff,
            }
        )

    unmatched_budget = [idx for idx in range(len(budget_rows)) if idx not in used_budget]

    summary = {
        "budget_rows": len(budget_rows),
        "quantity_rows": len(quantity_rows),
        "matched_rows": len(matches),
        "unmatched_budget": len(unmatched_budget),
        "unmatched_quantity": len(unmatched_quantity),
    }
    return {
        "summary": summary,
        "matches": matches,
        "unmatched_budget": unmatched_budget,
        "unmatched_quantity": unmatched_quantity,
    }


def compare_rows(
    budget_rows: list[dict[str, Any]],
    quantity_rows: list[dict[str, Any]],
    synonym_map: dict[str, str] | None = None,
    threshold: int = 88,
    compare_mode: str = "balanced",
) -> dict:
    mode = (compare_mode or "balanced").strip().lower()
    if mode == "classic":
        return compare_rows_classic(
            budget_rows=budget_rows,
            quantity_rows=quantity_rows,
            synonym_map=synonym_map,
            threshold=threshold,
        )
    if mode == "strict":
        return compare_rows_token(
            budget_rows=budget_rows,
            quantity_rows=quantity_rows,
            synonym_map=synonym_map,
            threshold=threshold,
            allow_fuzzy=False,
            allow_global_fuzzy=False,
            min_overlap=0.8,
            min_common_tokens=3,
        )
    if mode == "lenient":
        return compare_rows_token(
            budget_rows=budget_rows,
            quantity_rows=quantity_rows,
            synonym_map=synonym_map,
            threshold=max(80, threshold - 6),
            allow_fuzzy=True,
            allow_global_fuzzy=True,
            min_overlap=0.3,
            min_common_tokens=1,
            max_candidates=200,
        )
    return compare_rows_token(
        budget_rows=budget_rows,
        quantity_rows=quantity_rows,
        synonym_map=synonym_map,
        threshold=threshold,
        allow_fuzzy=True,
        allow_global_fuzzy=True,
        min_overlap=MIN_TOKEN_OVERLAP,
        min_common_tokens=MIN_COMMON_TOKENS,
        max_candidates=MAX_CANDIDATES,
    )


def build_report(
    budget_rows: list[dict[str, Any]],
    quantity_rows: list[dict[str, Any]],
    compare_result: dict,
    output_basename: str,
) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Merged"
    ws.append(
        [
            "數量_項次",
            "數量_分類",
            "數量_名稱",
            "數量_單位",
            "數量_數量",
            "數量_工作表",
            "預算_項次",
            "預算_參考碼",
            "預算_名稱",
            "預算_單位",
            "預算_數量",
            "名稱差異-預算獨有詞",
            "名稱差異-數量獨有詞",
            "名稱差異-預算獨有數字",
            "名稱差異-數量獨有數字",
            "單位一致",
            "數量差異",
            "匹配方法",
            "匹配分數",
        ]
    )

    for match in compare_result["matches"]:
        q = quantity_rows[match["quantity_index"]]
        b = budget_rows[match["budget_index"]]
        q_unit = _normalize_key(q.get("unit"))
        b_unit = _normalize_key(b.get("unit"))
        unit_match = "Y" if q_unit and b_unit and q_unit == b_unit else ""
        qty_delta = None
        if q.get("quantity") is not None and b.get("quantity") is not None:
            qty_delta = q.get("quantity") - b.get("quantity")
        ws.append(
            [
                q.get("item_no"),
                q.get("category"),
                q.get("description"),
                q.get("unit"),
                q.get("quantity"),
                q.get("sheet"),
                b.get("item_no"),
                b.get("ref_code"),
                b.get("description"),
                b.get("unit"),
                b.get("quantity"),
                _summarize_list(match.get("name_diff", {}).get("tokens_only_budget", [])),
                _summarize_list(match.get("name_diff", {}).get("tokens_only_quantity", [])),
                _summarize_list(match.get("name_diff", {}).get("numbers_only_budget", [])),
                _summarize_list(match.get("name_diff", {}).get("numbers_only_quantity", [])),
                unit_match,
                qty_delta,
                match["method"],
                match["score"],
            ]
        )

    wb.create_sheet("OnlyBudget")
    ws_budget = wb["OnlyBudget"]
    ws_budget.append(["ItemNo", "RefCode", "Description", "Unit", "Quantity"])
    for idx in compare_result["unmatched_budget"]:
        row = budget_rows[idx]
        ws_budget.append(
            [
                row.get("item_no"),
                row.get("ref_code"),
                row.get("description"),
                row.get("unit"),
                row.get("quantity"),
            ]
        )

    wb.create_sheet("OnlyQuantity")
    ws_qty = wb["OnlyQuantity"]
    ws_qty.append(["ItemNo", "Category", "Description", "Unit", "Quantity", "Sheet"])
    for idx in compare_result["unmatched_quantity"]:
        row = quantity_rows[idx]
        ws_qty.append(
            [
                row.get("item_no"),
                row.get("category"),
                row.get("description"),
                row.get("unit"),
                row.get("quantity"),
                row.get("sheet"),
            ]
        )

    date_path = datetime.now().strftime("%Y/%m/%d")
    report_dir = Path(settings.MEDIA_ROOT) / "design_consistency" / "reports" / date_path
    report_dir.mkdir(parents=True, exist_ok=True)
    output_path = report_dir / output_basename
    wb.save(output_path)
    return str(output_path.relative_to(settings.MEDIA_ROOT))


def run_compare(
    budget_file_path: str,
    quantity_file_path: str,
    synonym_map: dict[str, str] | None = None,
    threshold: int = 88,
    compare_mode: str = "balanced",
    budget_filter: str | None = None,
) -> tuple[dict, str, list[dict[str, Any]]]:
    budget_rows_raw = extract_budget_rows(budget_file_path)
    budget_rows, budget_stats = _filter_budget_rows(budget_rows_raw, budget_filter)
    quantity_rows = extract_quantity_rows(quantity_file_path)

    compare_result = compare_rows(
        budget_rows=budget_rows,
        quantity_rows=quantity_rows,
        synonym_map=synonym_map,
        threshold=threshold,
        compare_mode=compare_mode,
    )
    unit_mismatch = 0
    qty_mismatch = 0
    name_variance = 0
    matched_consistent = 0
    matched_items: list[dict[str, Any]] = []
    issue_items: list[dict[str, Any]] = []
    for match in compare_result["matches"]:
        q = quantity_rows[match["quantity_index"]]
        b = budget_rows[match["budget_index"]]
        name_diff = match.get("name_diff") or {}
        unit_match = None
        if _normalize_key(q.get("unit")) and _normalize_key(b.get("unit")):
            unit_match = _normalize_key(q.get("unit")) == _normalize_key(b.get("unit"))
        qty_delta = None
        qty_match = None
        if q.get("quantity") is not None and b.get("quantity") is not None:
            qty_delta = q.get("quantity") - b.get("quantity")
            qty_match = qty_delta == 0
        low_confidence = match["method"] in {"fuzzy", "global_fuzzy"} and match["score"] < 92
        is_consistent = (
            not name_diff.get("has_diff")
            and unit_match is not False
            and qty_match is not False
            and not low_confidence
        )
        if is_consistent:
            matched_consistent += 1
        matched_items.append(
            {
                "quantity": {
                    "item_no": q.get("item_no"),
                    "category": q.get("category"),
                    "description": q.get("description"),
                    "unit": q.get("unit"),
                    "quantity": q.get("quantity"),
                    "sheet": q.get("sheet"),
                    "row": q.get("row"),
                },
                "budget": {
                    "item_no": b.get("item_no"),
                    "ref_code": b.get("ref_code"),
                    "description": b.get("description"),
                    "unit": b.get("unit"),
                    "quantity": b.get("quantity"),
                },
                "method": match.get("method"),
                "score": match.get("score"),
                "unit_match": unit_match,
                "qty_delta": qty_delta,
                "qty_match": qty_match,
                "name_diff": name_diff,
                "consistent": is_consistent,
            }
        )
        if name_diff.get("has_diff"):
            name_variance += 1
            issue_items.append(
                {
                    "type": "name_variance",
                    "summary": f"{q.get('description')} <> {b.get('description')}",
                    "payload": {
                        "quantity": q,
                        "budget": b,
                        "score": match["score"],
                        "method": match["method"],
                        "name_diff": name_diff,
                    },
                }
            )
        if _normalize_key(q.get("unit")) and _normalize_key(b.get("unit")):
            if _normalize_key(q.get("unit")) != _normalize_key(b.get("unit")):
                unit_mismatch += 1
                issue_items.append(
                    {
                        "type": "unit_mismatch",
                        "summary": f"{q.get('description')} ⟂ {b.get('description')}",
                        "payload": {
                            "quantity": q,
                            "budget": b,
                            "score": match["score"],
                            "method": match["method"],
                        },
                    }
                )
        if q.get("quantity") is not None and b.get("quantity") is not None:
            if q.get("quantity") != b.get("quantity"):
                qty_mismatch += 1
                issue_items.append(
                    {
                        "type": "qty_mismatch",
                        "summary": f"{q.get('description')} Δ {b.get('description')}",
                        "payload": {
                            "quantity": q,
                            "budget": b,
                            "score": match["score"],
                            "method": match["method"],
                        },
                    }
                )
        if match["method"] in {"fuzzy", "global_fuzzy"} and match["score"] < 92:
            issue_items.append(
                {
                    "type": "low_confidence",
                    "summary": f"{q.get('description')} ≈ {b.get('description')}",
                    "payload": {
                        "quantity": q,
                        "budget": b,
                        "score": match["score"],
                        "method": match["method"],
                    },
                }
            )
    term_budget = build_term_stats([row["description"] for row in budget_rows], synonym_map)
    term_quantity = build_term_stats([row["description"] for row in quantity_rows], synonym_map)
    result = {
        "summary": {
            **compare_result["summary"],
            "name_variance": name_variance,
            "unit_mismatch": unit_mismatch,
            "qty_mismatch": qty_mismatch,
            "matched_consistent": matched_consistent,
            "matched_inconsistent": len(compare_result["matches"]) - matched_consistent,
            "compare_mode": compare_mode,
            "budget_filter": budget_filter or "",
            "budget_rows_total": budget_stats["total"],
            "budget_rows_skipped": budget_stats["skipped"],
        },
        "budget_terms": term_budget,
        "quantity_terms": term_quantity,
        "only_in_budget": [
            budget_rows[idx]["description"] for idx in compare_result["unmatched_budget"][:40]
        ],
        "only_in_quantity": [
            quantity_rows[idx]["description"] for idx in compare_result["unmatched_quantity"][:40]
        ],
        "matched_items": matched_items,
    }

    output_basename = f"comparison_{Path(budget_file_path).stem}_{Path(quantity_file_path).stem}.xlsx"
    report_rel_path = build_report(budget_rows, quantity_rows, compare_result, output_basename)
    for idx in compare_result["unmatched_budget"]:
        row = budget_rows[idx]
        issue_items.append(
            {
                "type": "unmatched_budget",
                "summary": row.get("description") or row.get("item_no") or "Budget only",
                "payload": {"budget": row},
            }
        )
    for idx in compare_result["unmatched_quantity"]:
        row = quantity_rows[idx]
        issue_items.append(
            {
                "type": "unmatched_quantity",
                "summary": row.get("description") or row.get("item_no") or "Quantity only",
                "payload": {"quantity": row},
            }
        )
    return result, report_rel_path, issue_items


# ─────────────────────────────────────────────────────────────────────────────
# File Structure Preview Functions
# ─────────────────────────────────────────────────────────────────────────────


def parse_xls_structure(file_path: str, max_preview_rows: int = 5) -> dict[str, Any]:
    """
    Parse an XLS/XLSX file and return its structure for preview.
    
    Returns:
        dict with keys:
        - filename: original filename
        - total_sheets: number of sheets
        - sheets: list of sheet info dicts
        - summary: overall summary stats
    """
    ext = Path(file_path).suffix.lower()
    filename = Path(file_path).name
    sheets_info: list[dict[str, Any]] = []
    
    if ext == ".xls":
        try:
            import xlrd
        except ImportError as exc:
            raise RuntimeError("Missing xlrd for .xls parsing.") from exc
        
        wb = xlrd.open_workbook(file_path)
        for idx, sheet_name in enumerate(wb.sheet_names()):
            sheet = wb.sheet_by_index(idx)
            nrows = sheet.nrows
            ncols = sheet.ncols
            
            # Get sample rows
            preview_rows: list[list[str]] = []
            for r in range(min(max_preview_rows, nrows)):
                row_values = []
                for c in range(min(10, ncols)):  # Limit to 10 columns for preview
                    val = sheet.cell_value(r, c)
                    row_values.append(_normalize_text(val)[:50])  # Truncate long values
                preview_rows.append(row_values)
            
            # Detect header row
            try:
                import pandas as pd
                frame = pd.read_excel(file_path, sheet_name=sheet_name, header=None, engine="xlrd")
                values_list = frame.fillna("").values.tolist()
                header_row, header_map = _detect_header_row_values(values_list)
                has_structure = "description" in header_map and "quantity" in header_map
                estimated_items = 0
                if has_structure:
                    # Count rows with both description and quantity
                    desc_col = header_map.get("description", 0) - 1
                    qty_col = header_map.get("quantity", 0) - 1
                    for r in range(header_row, len(values_list)):
                        row_vals = values_list[r]
                        if len(row_vals) > max(desc_col, qty_col):
                            desc = _normalize_text(row_vals[desc_col]) if desc_col >= 0 else ""
                            qty = _parse_number(row_vals[qty_col]) if qty_col >= 0 else None
                            if desc and qty is not None:
                                estimated_items += 1
            except Exception:
                header_row = 0
                header_map = {}
                has_structure = False
                estimated_items = 0
            
            sheets_info.append({
                "index": idx,
                "name": sheet_name,
                "rows": nrows,
                "cols": ncols,
                "header_row": header_row,
                "has_structure": has_structure,
                "detected_columns": list(header_map.keys()),
                "estimated_items": estimated_items,
                "preview": preview_rows,
            })
    else:
        # XLSX format using openpyxl
        wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
        for idx, ws in enumerate(wb.worksheets):
            nrows = ws.max_row or 0
            ncols = ws.max_column or 0
            
            # Get sample rows
            preview_rows: list[list[str]] = []
            for r, row in enumerate(ws.iter_rows(max_row=max_preview_rows, max_col=10)):
                row_values = [_normalize_text(cell.value)[:50] for cell in row]
                preview_rows.append(row_values)
            
            # Need to reload for header detection (read_only mode limitation)
            try:
                wb_detect = openpyxl.load_workbook(file_path, data_only=True)
                ws_detect = wb_detect[ws.title]
                header_row, header_map = _detect_header_row(ws_detect)
                has_structure = "description" in header_map and "quantity" in header_map
                estimated_items = 0
                if has_structure:
                    for r in range(header_row + 1, (ws_detect.max_row or 0) + 1):
                        desc = _normalize_text(ws_detect.cell(row=r, column=header_map.get("description", 1)).value)
                        qty = _parse_number(ws_detect.cell(row=r, column=header_map.get("quantity", 1)).value)
                        if desc and qty is not None:
                            estimated_items += 1
                wb_detect.close()
            except Exception:
                header_row = 0
                header_map = {}
                has_structure = False
                estimated_items = 0
            
            sheets_info.append({
                "index": idx,
                "name": ws.title,
                "rows": nrows,
                "cols": ncols,
                "header_row": header_row,
                "has_structure": has_structure,
                "detected_columns": list(header_map.keys()),
                "estimated_items": estimated_items,
                "preview": preview_rows,
            })
        wb.close()
    
    # Build summary
    total_rows = sum(s["rows"] for s in sheets_info)
    total_items = sum(s["estimated_items"] for s in sheets_info)
    parseable_sheets = sum(1 for s in sheets_info if s["has_structure"])
    
    return {
        "filename": filename,
        "file_type": "xls" if ext == ".xls" else "xlsx",
        "total_sheets": len(sheets_info),
        "sheets": sheets_info,
        "summary": {
            "total_rows": total_rows,
            "total_estimated_items": total_items,
            "parseable_sheets": parseable_sheets,
            "unparseable_sheets": len(sheets_info) - parseable_sheets,
        },
    }


def _walk_xml_items_with_depth(elem, depth: int = 0) -> list[dict[str, Any]]:
    """Walk XML items and record their depth level."""
    row = _build_budget_row(elem)
    row["depth"] = depth
    rows = [row]
    for child in elem:
        if child.tag.endswith("PayItem") or child.tag.endswith("WorkItem"):
            rows.extend(_walk_xml_items_with_depth(child, depth + 1))
    return rows


def parse_xml_structure(file_path: str, max_children_preview: int = 10) -> dict[str, Any]:
    """
    Parse an XML budget file and return its structure for preview.
    
    Returns:
        dict with keys:
        - filename: original filename
        - contract_info: contract/tender information
        - hierarchy: hierarchical structure of PayItems
        - item_type_distribution: count by item_kind
        - summary: overall summary stats
    """
    if etree is None:
        raise RuntimeError("Missing lxml for XML parsing.")
    
    filename = Path(file_path).name
    parser = etree.XMLParser(recover=True, huge_tree=True)
    tree = etree.parse(file_path, parser)
    root = tree.getroot()
    
    # Extract contract information
    contract_info: dict[str, Any] = {}
    tender_info = root.find(".//ns:TenderInformation", ETENDER_NS)
    if tender_info is not None:
        contract_info["contract_no"] = tender_info.get("contractNo", "")
        # Get procuring entity
        entity = tender_info.find('ns:ProcuringEntity[@language="zh-TW"]', ETENDER_NS)
        if entity is not None and entity.text:
            contract_info["procuring_entity"] = _normalize_text(entity.text)
        # Get contract title
        title = tender_info.find('ns:ContractTitle[@language="zh-TW"]', ETENDER_NS)
        if title is not None and title.text:
            contract_info["contract_title"] = _normalize_text(title.text)
        # Get contract location
        location = tender_info.find('ns:ContractLocation', ETENDER_NS)
        if location is not None and location.text:
            contract_info["contract_location"] = _normalize_text(location.text)
    
    # Build hierarchy from DetailList
    hierarchy: list[dict[str, Any]] = []
    all_items: list[dict[str, Any]] = []
    
    detail_list = root.find(".//ns:DetailList", ETENDER_NS)
    if detail_list is not None:
        for elem in detail_list.findall("ns:PayItem", ETENDER_NS):
            items = _walk_xml_items_with_depth(elem, depth=0)
            all_items.extend(items)
            
            # Build hierarchy entry for top-level items
            row = _build_budget_row(elem)
            child_count = len([e for e in elem if e.tag.endswith("PayItem") or e.tag.endswith("WorkItem")])
            total_descendants = len(items) - 1  # Exclude self
            
            # Get first few children for preview
            children_preview: list[dict[str, str]] = []
            for child in list(elem)[:max_children_preview]:
                if child.tag.endswith("PayItem") or child.tag.endswith("WorkItem"):
                    child_row = _build_budget_row(child)
                    children_preview.append({
                        "item_no": child_row.get("item_no", ""),
                        "description": (child_row.get("description", "")[:60] + "...") if len(child_row.get("description", "")) > 60 else child_row.get("description", ""),
                        "item_kind": child_row.get("item_kind", ""),
                    })
            
            hierarchy.append({
                "item_no": row.get("item_no", ""),
                "description": row.get("description", ""),
                "item_kind": row.get("item_kind", ""),
                "direct_children": child_count,
                "total_descendants": total_descendants,
                "children_preview": children_preview,
            })
    
    # Also check CostBreakdownList
    cbl = root.find(".//ns:CostBreakdownList", ETENDER_NS)
    if cbl is not None:
        for elem in cbl.findall("ns:WorkItem", ETENDER_NS):
            items = _walk_xml_items_with_depth(elem, depth=0)
            all_items.extend(items)
    
    # Calculate item type distribution
    item_kind_counter = Counter()
    depth_counter = Counter()
    items_with_quantity = 0
    
    for item in all_items:
        kind = item.get("item_kind", "unknown") or "unknown"
        item_kind_counter[kind] += 1
        depth_counter[item.get("depth", 0)] += 1
        if item.get("quantity") is not None:
            items_with_quantity += 1
    
    kind_labels = {
        "mainitem": "主項 (MainItem)",
        "analysis": "分析 (Analysis)",
        "subtotal": "小計 (Subtotal)",
        "general": "一般 (General)",
        "unknown": "未知 (Unknown)",
    }
    
    item_type_distribution = [
        {
            "kind": kind,
            "label": kind_labels.get(kind, kind),
            "count": count,
        }
        for kind, count in sorted(item_kind_counter.items(), key=lambda x: -x[1])
    ]
    
    depth_distribution = [
        {"depth": depth, "count": count}
        for depth, count in sorted(depth_counter.items())
    ]
    
    max_depth = max(depth_counter.keys()) if depth_counter else 0
    
    return {
        "filename": filename,
        "file_type": "xml",
        "contract_info": contract_info,
        "hierarchy": hierarchy,
        "item_type_distribution": item_type_distribution,
        "depth_distribution": depth_distribution,
        "summary": {
            "total_items": len(all_items),
            "items_with_quantity": items_with_quantity,
            "items_without_quantity": len(all_items) - items_with_quantity,
            "top_level_items": len(hierarchy),
            "max_depth": max_depth,
            "unique_item_kinds": len(item_kind_counter),
        },
    }
